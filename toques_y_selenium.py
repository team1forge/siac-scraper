"""
toques_y_selenium.py
La VPN la conecta una persona a mano ANTES de ejecutar el script (el script no la toca).
1) Abre el portal de distribuidores y hace clic en la opción SIAC Único.
2) Por cada RUC del Excel: busca, extrae los datos al CSV, cierra todo y reabre SIAC Único.

Uso:
  python toques_y_selenium.py                       # con la VPN ya conectada: portal -> SIAC -> lote
  python toques_y_selenium.py --guardar-password    # 1ra vez: guarda la contraseña de forma segura
  python toques_y_selenium.py --verificar-password  # prueba la contraseña guardada contra el portal
  python toques_y_selenium.py --metricas            # resumen del CSV: OK / falta de datos / bloqueo / fallo
"""
import argparse
import base64
import csv
import logging
import re
import sys
import time
from pathlib import Path
from urllib.parse import unquote

import keyring
from pywinauto import Desktop
from pywinauto.keyboard import send_keys
from selenium import webdriver
from selenium.common.exceptions import UnexpectedAlertPresentException
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

# ============================ CONFIG ============================
# Login: la contraseña se guarda en el Administrador de credenciales de Windows
# (python toques_y_selenium.py --guardar-password), NUNCA en este archivo.
USUARIO_VPN = "D79983538"
KEYRING_SERVICIO = "ivanti_vpn"   # no cambiar: ahí ya está guardada la contraseña
PAUSA_TECLA = 0.15            # segundos entre tecla y tecla (súbelo si el diálogo se pierde caracteres)
ESPERA_FOCO = 3               # segundos antes de empezar a teclear (que el diálogo tome el foco)

URL_SELENIUM = "http://portaldistribuidores.tim.com.pe:1616/SitePages/inicio.aspx"
TEXTO_OPCION = "SIAC Unico"   # opción a buscar y clicar dentro del portal
TIPO_BUSQUEDA = "Documento Identidad"   # opción del desplegable de búsqueda en SIAC Único
VALOR_BUSQUEDA = ""           # un solo valor. Vacío = itera el Excel de abajo
# Excel con los valores a buscar (en la misma carpeta que este script)
ARCHIVO_VALORES = "80k.xlsx"
COLUMNA_VALORES = None        # None = autodetecta (RUC/documento/DNI); o el nombre o índice de la columna
HOJA_VALORES = 0              # índice o nombre de la hoja
PAUSA_ENTRE_BUSQUEDAS = 3     # segundos entre una búsqueda y la siguiente
# Popup "Consulta de Productos" que sale al presionar la lupa
TIMEOUT_POPUP = 60            # segundos esperando que aparezca el popup
REGEX_POPUP = re.compile(r"popup|consulta de productos", re.I)
# Popup "Consulta de Clientes": sale cuando el número tiene varios clientes (RUC, Pasaporte...).
# Se marca la fila cuyo Tipo Documento es TIPO_DOC_CLIENTE y se presiona Seleccionar
REGEX_CONSULTA_CLIENTES = re.compile(r"consulta de clientes", re.I)
TIPO_DOC_CLIENTE = "RUC"
# Popup de "Alerta" (ej. "Usted no está autorizado para consultar este tipo de producto"):
# se acepta, el RUC se guarda con datos vacíos y se pasa al siguiente
REGEX_ALERTA = re.compile(r"no est[aá] autorizad|^\s*alerta\s*$", re.I | re.M)
TIMEOUT_CIERRE_ALERTA = 15    # segundos esperando que la alerta se cierre tras el clic en Aceptar
TIMEOUT_DATOS = 30            # segundos esperando que el popup termine de pintar sus tablas
ESPERA_DATOS = 2              # segundos entre lectura y lectura del popup
ESPERA_POPUP_VACIO = 5        # segundos con las tablas sin filas para dar el popup por vacío
CLIC_EN_TOTAL_ACTIVAS = True  # tras leer los totales, clic en el número de activas (abre otro popup)
TIMEOUT_BOTON_TOTAL = 15      # segundos esperando que el recuadro Total Activas / No Activas esté clicable
ESPERA_POPUP_DETALLE = 10     # segundos esperando el popup de líneas tras cada intento de clic
# Título del detalle de líneas (puede abrirse en ventana nueva o dentro del mismo popup)
REGEX_DETALLE = re.compile(r"l[ií]neas asociadas", re.I)
ESPERA_TRAS_NUMERO = 5        # segundos en SIAC después de clicar el primer número del detalle
ARCHIVO_CSV = "resultado_rucs.csv"    # columnas: ruc + CAMPOS_CLIENTE
TIMEOUT_WEB = 40              # segundos esperando la carga del portal / la opción
TIMEOUT_VENTANA_SIAC = 30     # segundos esperando que SIAC abra su ventana nueva
REGEX_VENTANA_SIAC = re.compile(r"siac", re.I)   # título o URL de la ventana de SIAC
# Dominios donde Edge puede usar tu sesión de Windows (evita el popup de usuario/contraseña)
AUTH_ALLOWLIST = "*tim.com.pe,portaldistribuidores.tim.com.pe"
# Credenciales del portal: por defecto reutiliza las mismas de la VPN
REUSAR_CREDENCIALES_VPN = True        # False si el portal usa otro usuario/contraseña
USUARIO_PORTAL = USUARIO_VPN          # ej. "TIM\\D79983538" si el portal pide dominio
KEYRING_SERVICIO_PORTAL = "portal_distribuidores"
REGEX_DIALOGO_WEB = re.compile(r"(Seguridad de Windows|Windows Security|Iniciar sesi[oó]n|Sign in)", re.I)
NAVEGADOR = "chrome"          # "chrome" o "edge"
# Perfil dedicado para conservar cookies/sesión (SSO) entre ejecuciones. None = perfil temporal
PERFIL_DIR = Path(f"./perfil_{NAVEGADOR}").resolve()
MANTENER_ABIERTO = False      # True = deja el navegador abierto al terminar el script
# Si el navegador se cierra a mitad del proceso, se abre de nuevo y se retoma desde el CSV
MAX_REINICIOS_SIN_AVANCE = 3  # reinicios seguidos sin guardar ningún RUC nuevo antes de rendirse
MAX_CAIDAS_POR_RUC = 2        # si el navegador se cae N veces con el mismo RUC, lo guardo FALLO y sigo
# ================================================================

logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s")
log = logging.getLogger("bandeja")


# ------------------------ credenciales ------------------------
def _escritorio():
    return Desktop(backend="uia")


def _escapar_teclas(texto):
    """Escapa los caracteres especiales de send_keys (+ ^ % ~ ( ) { } [ ])."""
    return "".join("{%s}" % c if c in "+^%~(){}[]" else c for c in texto)


def guardar_password():
    import getpass
    pwd = getpass.getpass(f"Contraseña para {USUARIO_VPN} (no se mostrará): ")
    if not pwd:
        print("Contraseña vacía; no se guardó nada.")
        return
    keyring.set_password(KEYRING_SERVICIO, USUARIO_VPN, pwd)
    print("Guardada en el Administrador de credenciales de Windows.")


# -------------------------- Selenium --------------------------
_driver_activo = None                     # el navegador abierto por abrir_selenium()


class NavegadorCerrado(RuntimeError):
    """El navegador (o el driver) se cerró: hay que abrirlo de nuevo y retomar."""


def _driver_vivo(driver):
    """True si el navegador sigue abierto y responde."""
    if driver is None:
        return False
    try:
        return bool(driver.window_handles)
    except UnexpectedAlertPresentException:
        return True                           # hay un alert() abierto, pero el navegador vive
    except Exception:
        return False                          # sesión inválida, chromedriver muerto, sin ventanas...


def abrir_selenium(url=URL_SELENIUM, navegar=True):
    if NAVEGADOR == "edge":
        opts = webdriver.EdgeOptions()
    else:
        opts = webdriver.ChromeOptions()

    opts.add_argument("--start-maximized")
    # 'none': get() no espera la carga, así podemos atender el popup de autenticación
    opts.set_capability("pageLoadStrategy", "none")
    if AUTH_ALLOWLIST:
        # Autenticación integrada de Windows: Edge responde solo al reto NTLM/Kerberos
        opts.add_argument(f"--auth-server-allowlist={AUTH_ALLOWLIST}")
        opts.add_argument(f"--auth-negotiate-delegate-allowlist={AUTH_ALLOWLIST}")
    if PERFIL_DIR:
        opts.add_argument(f"--user-data-dir={PERFIL_DIR}")
    if MANTENER_ABIERTO:
        opts.add_experimental_option("detach", True)

    # Selenium 4.6+ descarga el driver solo (Selenium Manager)
    global _driver_activo
    driver = webdriver.Edge(options=opts) if NAVEGADOR == "edge" else webdriver.Chrome(options=opts)
    _driver_activo = driver               # para cerrarlo al final aunque el flujo falle
    if navegar:
        driver.get(url)
        WebDriverWait(driver, TIMEOUT_WEB).until(EC.presence_of_element_located((By.TAG_NAME, "body")))
        log.info("Selenium abierto en: %s", driver.current_url)
    return driver


def _buscar_en_frames(driver, xpath, timeout=TIMEOUT_WEB):
    """Busca el xpath en la página y dentro de cada iframe (SharePoint suele usarlos)."""
    fin = time.time() + timeout
    while time.time() < fin:
        driver.switch_to.default_content()
        contextos = [None] + driver.find_elements(By.TAG_NAME, "iframe")
        for frame in contextos:
            try:
                driver.switch_to.default_content()
                if frame is not None:
                    driver.switch_to.frame(frame)
                elementos = [el for el in driver.find_elements(By.XPATH, xpath) if el.is_displayed()]
                if elementos:
                    return elementos[0]
            except Exception:
                continue
        time.sleep(1)
    driver.switch_to.default_content()
    return None


_MAYUS = "ABCDEFGHIJKLMNOPQRSTUVWXYZÁÉÍÓÚÜÑ"
_MINUS = "abcdefghijklmnopqrstuvwxyzaeiouun"


def _normaliza(texto):
    return texto.lower().translate(str.maketrans("áéíóúüñ", "aeiouun"))


def _xpath_texto(tags, texto):
    """XPath que busca 'texto' en esos tags, sin distinguir mayúsculas ni tildes."""
    obj = _normaliza(texto)
    tr = f"translate(normalize-space(.),'{_MAYUS}','{_MINUS}')"
    return " | ".join(f"//{t}[contains({tr},'{obj}')]" for t in tags)


# Textos posibles del botón del desplegable (el tipo de búsqueda activo)
OPCIONES_BUSQUEDA = ("Teléfono", "Telefono", "Cuenta", "Customer Id", "ICCID", "Nombres",
                     "Documento Identidad", "Número de Recibo", "Razón Social")


def _buscar_toggle(driver):
    """El botón del desplegable de tipo de búsqueda (no el menú del usuario)."""
    # 0) Un botón/desplegable cuyo texto ES un tipo de búsqueda (en la ficha del cliente
    #    hay muchos textos con "Cuenta"; así no confunde una etiqueta con el botón)
    opciones = {_normaliza(o) for o in OPCIONES_BUSQUEDA}
    for el in driver.find_elements(By.XPATH, "//button | //*[contains(@class,'dropdown-toggle')]"
                                             " | //*[@data-toggle='dropdown']"):
        try:
            if el.is_displayed() and _normaliza(" ".join(el.text.split())) in opciones:
                return el
        except Exception:
            continue
    # 1) Por su texto actual: es el único que muestra un tipo de búsqueda
    for opcion in OPCIONES_BUSQUEDA:
        xp = _xpath_texto(("button", "a", "span", "div"), opcion)
        for el in driver.find_elements(By.XPATH, xp):
            try:
                if el.is_displayed() and len(el.text.strip()) < 25:
                    return el
            except Exception:
                continue
    # 2) Respaldo: un dropdown-toggle visible con texto corto
    for el in driver.find_elements(By.XPATH, "//*[contains(@class,'dropdown-toggle')]"
                                             " | //*[@data-toggle='dropdown']"):
        try:
            if el.is_displayed() and 0 < len(el.text.strip()) < 25:
                return el
        except Exception:
            continue
    return None


def seleccionar_tipo_busqueda(driver, texto=TIPO_BUSQUEDA):
    """En SIAC Único: abre el desplegable de tipo de búsqueda y elige la opción indicada."""
    log.info("[8/8] SIAC: buscando el botón del desplegable...")
    toggle = _buscar_toggle(driver)
    if toggle is None:
        raise RuntimeError("No encontré el desplegable de tipo de búsqueda en SIAC.")

    actual = toggle.text.strip()
    log.info("[8/8] SIAC: botón encontrado -> %r. Haciendo clic para desplegarlo...", actual)
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", toggle)
    time.sleep(0.5)
    try:
        toggle.click()
    except Exception:
        driver.execute_script("arguments[0].click();", toggle)
    time.sleep(1.5)
    log.info("[8/8] SIAC: desplegable abierto. Buscando la opción %r...", texto)

    # La opción solo existe/es visible con el menú abierto
    opcion = None
    for el in driver.find_elements(By.XPATH, _xpath_texto(("a", "li", "span", "button"), texto)):
        try:
            if el.is_displayed():
                opcion = el
                break
        except Exception:
            continue

    if opcion is None:
        visibles = sorted({el.text.strip() for el in driver.find_elements(By.XPATH, "//li | //a")
                           if el.text.strip() and len(el.text.strip()) < 40 and el.is_displayed()})[:30]
        raise RuntimeError(
            f"Abrí el desplegable pero no encontré '{texto}'. Opciones visibles: {', '.join(visibles)}"
        )

    log.info("[8/8] SIAC: opción encontrada -> %r. Haciendo clic...", opcion.text.strip())
    try:
        opcion.click()
    except Exception:
        driver.execute_script("arguments[0].click();", opcion)
    time.sleep(1)
    toggle = _buscar_toggle(driver)
    nuevo = toggle.text.strip() if toggle is not None else ""
    if _normaliza(texto) in _normaliza(nuevo):
        log.info("[8/8] LISTO: el tipo de búsqueda cambió de %r a %r", actual, nuevo)
    else:
        log.warning("[8/8] El botón sigue mostrando %r; puede que el clic no aplicara.", nuevo)
    return opcion


def buscar_en_siac(driver, valor=VALOR_BUSQUEDA):
    """Escribe el valor en la caja de búsqueda de SIAC y hace clic en la lupa."""
    if not valor:
        log.info("[8/8] SIAC: sin valor que buscar (VALOR_BUSQUEDA está vacío).")
        return False

    # La caja de texto de la barra roja (junto al desplegable)
    caja = None
    for el in driver.find_elements(By.XPATH, "//input[@type='text'] | //input[not(@type)]"):
        try:
            if el.is_displayed() and el.is_enabled():
                caja = el
                break
        except Exception:
            continue
    if caja is None:
        raise RuntimeError("No encontré la caja de búsqueda en SIAC.")

    log.info("[8/8] SIAC: caja de búsqueda encontrada (name=%r, id=%r).",
             caja.get_attribute("name"), caja.get_attribute("id"))
    log.info("[8/8] SIAC: escribiendo %r en el input...", valor)
    caja.clear()
    caja.send_keys(valor)
    time.sleep(0.5)

    escrito = caja.get_attribute("value")
    if escrito == valor:
        log.info("[8/8] SIAC: dato ingresado OK en el input -> %r", escrito)
    else:
        log.warning("[8/8] SIAC: el input quedó con %r en vez de %r", escrito, valor)

    # La lupa: botón con ícono de búsqueda al lado de la caja
    lupa = None
    for xp in ("//button[contains(@class,'search')]", "//*[contains(@class,'glyphicon-search')]/..",
               "//*[contains(@class,'fa-search')]/..", "//button[@type='submit']"):
        els = [e for e in driver.find_elements(By.XPATH, xp) if e.is_displayed()]
        if els:
            lupa = els[0]
            break

    if lupa is not None:
        log.info("[8/8] SIAC: botón de búsqueda encontrado; presionándolo...")
        try:
            lupa.click()
            log.info("[8/8] SIAC: BOTÓN PRESIONADO (clic normal).")
        except Exception as e:
            log.info("[8/8] SIAC: clic normal falló (%s); presionando por JavaScript...", type(e).__name__)
            driver.execute_script("arguments[0].click();", lupa)
            log.info("[8/8] SIAC: BOTÓN PRESIONADO (vía JavaScript).")
    else:
        log.info("[8/8] SIAC: no ubiqué el botón de la lupa; enviando ENTER en el input...")
        caja.send_keys(Keys.ENTER)
        log.info("[8/8] SIAC: ENTER enviado.")

    time.sleep(3)
    log.info("[8/8] SIAC: búsqueda enviada. URL actual: %s", driver.current_url)
    return True


def leer_valores(ruta=ARCHIVO_VALORES, columna=COLUMNA_VALORES, hoja=HOJA_VALORES):
    """Lee el Excel y devuelve los valores a buscar, sin vacíos ni duplicados."""
    from openpyxl import load_workbook

    archivo = Path(ruta)
    if not archivo.exists():                       # por si el nombre vino sin extensión
        candidatos = sorted(Path(".").glob(archivo.stem + ".xls*"))
        if not candidatos:
            raise RuntimeError(f"No encontré '{ruta}' en {Path('.').resolve()}")
        archivo = candidatos[0]

    wb = load_workbook(archivo, read_only=True, data_only=True)
    ws = wb[hoja] if isinstance(hoja, str) else wb.worksheets[hoja]
    filas = ws.iter_rows(values_only=True)
    cabecera = ["" if c is None else str(c).strip() for c in next(filas, ())]

    if isinstance(columna, int):
        idx = columna
    elif columna:
        idx = cabecera.index(columna)
    else:
        idx = next((i for i, h in enumerate(cabecera) if re.search(r"ruc|documento|dni", h, re.I)), None)
        if idx is None:
            idx = 0
            log.warning("[8/8] No identifiqué la columna; uso la primera. Encabezados: %s",
                        ", ".join(cabecera) or "(ninguno)")

    def limpiar(v):
        v = str(v).strip()
        return v[:-2] if v.endswith(".0") else v    # openpyxl devuelve los números como 20123456789.0

    valores, vistos = [], set()
    primera = cabecera[idx] if idx < len(cabecera) else ""
    if primera.replace(" ", "").isdigit() and len(primera) >= 8:
        valores.append(limpiar(primera))            # el archivo no tenía encabezados

    for fila in filas:
        if idx >= len(fila) or fila[idx] is None:
            continue
        v = limpiar(fila[idx])
        if v and v not in vistos:
            vistos.add(v)
            valores.append(v)
    wb.close()

    log.info("[8/8] %s: %d valores en la columna %r", archivo.name, len(valores),
             cabecera[idx] if idx < len(cabecera) else idx)
    return valores


def _valor_celda(celda):
    """El número puede venir dentro de un <input> o como texto plano."""
    inputs = celda.find_elements(By.TAG_NAME, "input")
    if inputs:
        return (inputs[0].get_attribute("value") or "").strip()
    return celda.text.strip()


def _a_entero(texto):
    digitos = re.sub(r"[^0-9]", "", texto or "")
    return int(digitos) if digitos else None


def esperar_popup(driver, handles_antes, timeout=TIMEOUT_POPUP):
    """Devuelve el handle del popup de resultados, o None si no apareció."""
    fin = time.time() + timeout
    nuevo = None
    while time.time() < fin:
        for h in driver.window_handles:
            if h in handles_antes:
                continue
            try:
                driver.switch_to.window(h)
                nuevo = h
                if _texto_alerta(driver):             # la ventana nueva es una alerta: no espero más
                    log.info("[8/8] Ventana nueva con alerta: %s", driver.current_url)
                    return h
                if REGEX_POPUP.search(driver.current_url) or REGEX_POPUP.search(driver.title or ""):
                    log.info("[8/8] Popup detectado: %s", driver.current_url)
                    return h
            except Exception:
                continue          # todavía abriendo o ya se cerró
        time.sleep(1)

    if nuevo and nuevo in driver.window_handles:
        log.info("[8/8] Ventana nueva sin el patrón esperado; la uso igual.")
        return nuevo
    return None


def _filas_con_datos(tabla):
    return [f for f in tabla.find_elements(By.XPATH, ".//tbody/tr") if f.find_elements(By.TAG_NAME, "td")]


def _tablas_servicios(driver):
    """Tablas de servicios del popup: [(encabezados, filas), ...].

    La tabla tiene scroll propio: a veces el encabezado y el cuerpo son dos <table>
    distintas (encabezado fijo arriba). En ese caso las filas salen de la tabla siguiente.
    """
    tablas = []
    for tabla in driver.find_elements(By.TAG_NAME, "table"):
        try:
            cabecera = [" ".join(th.text.split()) for th in tabla.find_elements(By.TAG_NAME, "th")]
            if not any("total activas" in _normaliza(h) for h in cabecera):
                continue
            filas = _filas_con_datos(tabla)
            if not filas:
                siguiente = tabla.find_elements(By.XPATH, "following::table[1]")
                # solo si es el cuerpo (sin encabezados visibles), no la tabla de otra sección
                if siguiente and not any(th.text.strip() for th in siguiente[0].find_elements(By.TAG_NAME, "th")):
                    filas = _filas_con_datos(siguiente[0])
            tablas.append((cabecera, filas))
        except Exception:
            continue                          # la tabla se repintó mientras la leía
    return tablas


REGEX_SIN_DATOS = re.compile(r"no existen datos|no hay datos|sin datos|ning[uú]n dato", re.I)


def _popup_sin_datos(driver):
    """True si la tabla de servicios POST-PAGO ya cargó y solo dice "No existen datos"."""
    tablas = _tablas_servicios(driver)
    if not tablas:
        return False
    try:
        for _, filas in tablas:
            textos = [" ".join(f.text.split()) for f in filas]
            if not textos or not all(REGEX_SIN_DATOS.search(t) for t in textos):
                return False
    except Exception:
        return False                          # se repintó mientras la leía
    return True


def _tablas_sin_filas(driver):
    """True si las tablas de servicios están cargadas pero sin ninguna fila con texto
    (pasa cuando el documento no tiene productos: solo se ven los encabezados)."""
    tablas = _tablas_servicios(driver)
    if not tablas:
        return False
    try:
        return all(not any(f.text.strip() for f in filas) for _, filas in tablas)
    except Exception:
        return False                          # se repintó mientras la leía


def _indice_activas(cabecera):
    return next((j for j, h in enumerate(cabecera)
                 if "total activas" in _normaliza(h) and "no activas" not in _normaliza(h)), None)


def _indice_no_activas(cabecera):
    return next((j for j, h in enumerate(cabecera) if "total no activas" in _normaliza(h)), None)


def _campo_celda(celda):
    """Lo que se clica: el enlace que envuelve el número (href="javascript:OpenDetail(...)"),
    si no el <input> habilitado con el número, si no la celda misma.
    (El input puede venir deshabilitado: un clic en él no llega al enlace.)"""
    enlaces = celda.find_elements(By.XPATH, ".//a | ./ancestor::a[1]")
    if enlaces:
        return enlaces[0]
    campos = celda.find_elements(By.TAG_NAME, "input")
    if campos and campos[0].is_enabled():
        return campos[0]
    return celda


def _campo_primera_fila(driver, clave, fila=0):
    """Vuelve a ubicar el recuadro "activas" o "no_activas" de esa fila (sin referencias viejas).
    Devuelve el elemento si está visible y habilitado, o None."""
    for cabecera, filas in _tablas_servicios(driver):
        i = _indice_activas(cabecera) if clave == "activas" else _indice_no_activas(cabecera)
        if i is None or len(filas) <= fila:
            continue
        try:
            celdas = filas[fila].find_elements(By.TAG_NAME, "td")
            if len(celdas) <= i:
                continue
            campo = _campo_celda(celdas[i])
            if campo.is_displayed():
                return campo
            if celdas[i].is_displayed():
                return celdas[i]              # el enlace no tiene tamaño propio: clic en la celda
        except Exception:
            continue                          # se repintó mientras lo buscaba
    return None


def esperar_boton_total(driver, clave, fila=0, timeout=TIMEOUT_BOTON_TOTAL):
    """Espera hasta que el recuadro de esa fila esté listo para clicarlo."""
    fin = time.time() + timeout
    while time.time() < fin:
        campo = _campo_primera_fila(driver, clave, fila)
        if campo is not None:
            return campo
        time.sleep(0.5)
    return None


# Qué eventos tiene el recuadro (para el log): HTML propio, de la celda y handlers de jQuery
_JS_DESCRIBIR = r"""
const el = arguments[0], td = el.closest('td'), tr = el.closest('tr');
const ev = x => { try { const d = window.jQuery && jQuery._data(x, 'events');
                        return d ? Object.keys(d).join(',') : ''; } catch (e) { return ''; } };
return {
  html: el.outerHTML.slice(0, 400),
  td: td ? td.outerHTML.slice(0, 400) : '',
  eventos_input: ev(el), eventos_td: td ? ev(td) : '', eventos_tr: tr ? ev(tr) : '',
};
"""

# Dispara la secuencia completa de eventos de mouse sobre el elemento (algunos sitios solo
# escuchan mousedown/mouseup o dblclick, no click)
_JS_EVENTOS_MOUSE = r"""
const el = arguments[0];
for (const tipo of ['mouseover', 'mousedown', 'focus', 'mouseup', 'click', 'dblclick']) {
  const e = tipo === 'focus' ? new FocusEvent('focus')
          : new MouseEvent(tipo, {bubbles: true, cancelable: true, view: window, detail: tipo === 'dblclick' ? 2 : 1});
  el.dispatchEvent(e);
}
"""


def _aparecio_detalle(driver, handles_antes, handle_popup):
    """Detecta el popup de líneas: una ventana nueva, o el mismo popup mostrando "Líneas asociadas"
    (a veces SIAC lo abre dentro de la misma ventana). Devuelve el handle o None."""
    nuevas = set(driver.window_handles) - set(handles_antes)
    if nuevas:
        return nuevas.pop()
    try:
        driver.switch_to.window(handle_popup)
        texto = driver.execute_script(_JS_TEXTO_VISIBLE) or ""
        if REGEX_DETALLE.search(texto):
            return handle_popup
    except Exception:
        pass
    return None


def clic_total_y_esperar_popup(driver, clave, handles_antes, fila=0):
    """Clic en el recuadro de esa fila probando varias formas de clic.

    Después de cada intento espera ESPERA_POPUP_DETALLE s a que aparezca el detalle;
    solo prueba la siguiente forma si no apareció (así no abre dos popups).
    Devuelve el handle donde quedó el detalle, o None."""
    handle_popup = driver.current_window_handle

    campo = esperar_boton_total(driver, clave, fila)
    if campo is None:
        log.warning("[8/8] No encontré el recuadro de la fila %d para clicarlo.", fila + 1)
        return None
    try:
        info = driver.execute_script(_JS_DESCRIBIR, campo)
        log.info("[8/8] Recuadro: %s", info["html"])
        log.info("[8/8] Celda:    %s", info["td"])
        log.info("[8/8] Eventos jQuery -> input: %r | td: %r | tr: %r",
                 info["eventos_input"], info["eventos_td"], info["eventos_tr"])
    except Exception:
        pass

    def celda(el):
        return el.find_element(By.XPATH, "./ancestor-or-self::td[1]")

    def ejecutar_enlace(el):
        """Corre el JavaScript del enlace (ej. OpenDetail('57268085','01','','A','MOVIL','18'))."""
        enlaces = el.find_elements(By.XPATH, "./ancestor-or-self::a[1] | .//a")
        href = (enlaces[0].get_attribute("href") or "") if enlaces else ""
        if not href.lower().startswith("javascript:"):
            raise RuntimeError("el recuadro no tiene enlace javascript:")
        codigo = unquote(href[len("javascript:"):])     # Chrome puede devolverlo con %20, %27...
        log.info("[8/8] Ejecutando el enlace del recuadro: %s", codigo)
        driver.execute_script(codigo)

    intentos = (
        ("enlace OpenDetail", ejecutar_enlace),
        ("clic normal", lambda el: el.click()),
        ("clic con el mouse", lambda el: ActionChains(driver).move_to_element(el).pause(0.3).click().perform()),
        ("foco + ENTER", lambda el: (el.click(), el.send_keys(Keys.ENTER))),
        ("eventos de mouse completos", lambda el: driver.execute_script(_JS_EVENTOS_MOUSE, el)),
        ("clic por JavaScript", lambda el: driver.execute_script("arguments[0].click();", el)),
        ("clic en la celda", lambda el: ActionChains(driver).move_to_element(celda(el)).pause(0.3).click().perform()),
        ("doble clic", lambda el: ActionChains(driver).double_click(el).perform()),
    )
    for nombre, clic in intentos:
        driver.switch_to.window(handle_popup)
        campo = esperar_boton_total(driver, clave, fila)
        if campo is None:
            log.warning("[8/8] El recuadro de la fila %d desapareció.", fila + 1)
            return _aparecio_detalle(driver, handles_antes, handle_popup)
        driver.execute_script("arguments[0].scrollIntoView({block:'center'});", campo)
        time.sleep(0.5)
        try:
            clic(campo)
        except Exception as e:
            log.info("[8/8] %s falló (%s); pruebo otra forma.", nombre, type(e).__name__)
            continue
        log.info("[8/8] %s hecho; espero el popup (hasta %s s)...", nombre.capitalize(), ESPERA_POPUP_DETALLE)
        fin = time.time() + ESPERA_POPUP_DETALLE
        while time.time() < fin:
            texto = _alerta_nativa(driver)
            if texto is not None:
                log.warning("[8/8] Salió una alerta al clicar: %r", texto)
            destino = _aparecio_detalle(driver, handles_antes, handle_popup)
            if destino:
                donde = "misma ventana" if destino == handle_popup else "ventana nueva"
                log.info("[8/8] Detalle abierto (%s) con %s.", donde, nombre)
                return destino
            time.sleep(0.5)
        log.info("[8/8] Con %s no se abrió el popup; pruebo otra forma.", nombre)
    return None


def esperar_primera_fila(driver, timeout=TIMEOUT_DATOS):
    """Espera a que la 1ra fila de productos tenga Producto y Total Activas y la imprime.

    Devuelve una lista con los totales de cada fila de la tabla: [(activas, no_activas), ...]
    (no_activas es 0 si la tabla no tiene esa columna); None si no cargó a tiempo."""
    log.info("[8/8] Esperando la primera fila de productos (hasta %s s)...", timeout)
    fin = time.time() + timeout
    while time.time() < fin:
        for cabecera, filas in _tablas_servicios(driver):
            i_act = _indice_activas(cabecera)
            if i_act is None or not filas:
                continue
            try:
                celdas = filas[0].find_elements(By.TAG_NAME, "td")
                if len(celdas) <= i_act:
                    continue
                valores = [_valor_celda(c) for c in celdas]
            except Exception:
                continue                      # se repintó; reintento
            if not (valores[0] and valores[i_act]):
                continue                      # todavía sin datos

            print("\nPrimera fila de productos:")
            for j, v in enumerate(valores):
                nombre = cabecera[j] if j < len(cabecera) and cabecera[j] else f"col {j}"
                print(f"  {nombre:<20} {v}")
            print()

            i_no = _indice_no_activas(cabecera)
            totales = []
            for fila in filas:
                try:
                    tds = fila.find_elements(By.TAG_NAME, "td")
                    act = _a_entero(_valor_celda(tds[i_act])) if i_act < len(tds) else None
                    no = _a_entero(_valor_celda(tds[i_no])) if i_no is not None and i_no < len(tds) else None
                except Exception:
                    act, no = None, None          # fila repintada a mitad de la lectura
                totales.append((act or 0, no or 0))
            return totales
        time.sleep(1)

    log.warning("[8/8] No detecté la primera fila en %s s.", timeout)
    return None


def _leer_tablas(driver):
    """Una pasada sobre las tablas del popup: [(clave_cuenta, activas, no_activas), ...]."""
    filas_datos = []
    for i, (cabecera, filas) in enumerate(_tablas_servicios(driver)):
        cabecera = [_normaliza(h) for h in cabecera]
        i_act = _indice_activas(cabecera)
        i_no = next((j for j, h in enumerate(cabecera) if "total no activas" in h), None)
        if i_act is None or i_no is None:
            continue                       # no es una tabla de servicios
        i_cta = next((j for j, h in enumerate(cabecera) if "cuenta" in h), None)

        for fila in filas:
            try:
                celdas = fila.find_elements(By.TAG_NAME, "td")
                if len(celdas) <= max(i_act, i_no):
                    continue
                act = _a_entero(_valor_celda(celdas[i_act]))
                no = _a_entero(_valor_celda(celdas[i_no]))
            except Exception:
                continue                   # fila repintada a mitad de la lectura
            if act is None and no is None:
                continue
            cuenta = _valor_celda(celdas[i_cta]) if i_cta is not None else str(len(filas_datos))
            filas_datos.append((f"{i}|{cuenta}", act or 0, no or 0))

    return filas_datos


# Texto visible de la página y de sus iframes (la alerta puede venir dentro de uno)
_JS_TEXTO_VISIBLE = r"""
let t = document.body ? document.body.innerText : '';
for (const f of document.querySelectorAll('iframe')) {
  try { t += '\n' + f.contentDocument.body.innerText; } catch (e) {}
}
return t;
"""


def _texto_alerta(driver):
    """El mensaje de la alerta si la ventana ACTUAL muestra una, o "" si no."""
    try:
        texto = driver.execute_script(_JS_TEXTO_VISIBLE) or ""
    except Exception:
        return ""
    if not REGEX_ALERTA.search(texto):
        return ""
    lineas = [l.strip() for l in texto.splitlines() if l.strip()]
    return next((l for l in lineas if REGEX_ALERTA.search(l) and _normaliza(l) != "alerta"), "Alerta")


def _alerta_nativa(driver):
    """Si hay un alert() de JavaScript abierto, lo acepta y devuelve su texto; si no, None."""
    try:
        alerta = driver.switch_to.alert
        texto = alerta.text or "Alerta"
        alerta.accept()
        return texto
    except Exception:
        return None


def buscar_ventana_alerta(driver, excluir=()):
    """Recorre las ventanas abiertas (menos las excluidas) buscando la alerta.
    Devuelve (handle, mensaje) y deja el driver en esa ventana, o (None, "")."""
    for h in driver.window_handles:
        if h in excluir:
            continue
        try:
            driver.switch_to.window(h)
        except Exception:
            continue                          # se cerró mientras recorría
        mensaje = _texto_alerta(driver)
        if mensaje:
            return h, mensaje
    return None, ""


def esperar_contenido_popup(driver, popup, excluir, timeout=TIMEOUT_DATOS):
    """Espera a que aparezcan las tablas del popup o una alerta (en cualquier ventana menos
    las de 'excluir', ej. SIAC y el portal).

    Devuelve ("datos", "", popup)
    | ("vacio", "", popup) si dice "No existen datos" o las tablas siguen sin filas ESPERA_POPUP_VACIO s
    | ("clientes", "", popup) si es la lista "Consulta de Clientes"
    | ("alerta", mensaje, handle_alerta) | ("nativa", mensaje, None)
    | (None, "", popup) si se agotó el tiempo.
    """
    fin = time.time() + timeout
    vacio_desde = None                        # desde cuándo las tablas se ven sin filas
    while time.time() < fin:
        texto = _alerta_nativa(driver)
        if texto is not None:
            return "nativa", texto, None

        h, mensaje = buscar_ventana_alerta(driver, excluir=excluir)
        if h:
            return "alerta", mensaje, h

        if popup in driver.window_handles:
            try:
                driver.switch_to.window(popup)
                if _leer_tablas(driver):
                    return "datos", "", popup
                if _popup_sin_datos(driver):
                    return "vacio", "", popup
                if REGEX_CONSULTA_CLIENTES.search(driver.execute_script(_JS_TEXTO_VISIBLE) or ""):
                    return "clientes", "", popup
                if _tablas_sin_filas(driver):
                    # puede estar cargando: solo lo doy por vacío si sigue así unos segundos
                    vacio_desde = vacio_desde or time.time()
                    if time.time() - vacio_desde >= ESPERA_POPUP_VACIO:
                        return "vacio", "", popup
                else:
                    vacio_desde = None
            except Exception:
                pass
        time.sleep(1)
    return None, "", popup


def _fila_cliente(driver, tipo=TIPO_DOC_CLIENTE):
    """En "Consulta de Clientes": la fila (<tr>) que tiene una celda exactamente igual a 'tipo'."""
    objetivo = _normaliza(tipo)
    for fila in driver.find_elements(By.XPATH, "//tbody/tr"):
        try:
            celdas = [_normaliza(" ".join(td.text.split())) for td in fila.find_elements(By.TAG_NAME, "td")]
            if objetivo in celdas and fila.is_displayed():
                return fila
        except Exception:
            continue                          # se repintó mientras la leía
    return None


def elegir_cliente_ruc(driver, popup, timeout=TIMEOUT_DATOS):
    """En "Consulta de Clientes": marca la fila con Tipo Documento = RUC y presiona Seleccionar.
    Devuelve el handle donde queda "Consulta de Productos" (la misma ventana u otra nueva),
    o None si no hay fila con RUC o no se pudo seleccionar."""
    driver.switch_to.window(popup)
    fin = time.time() + timeout
    fila = None
    while time.time() < fin and fila is None:
        fila = _fila_cliente(driver)
        if fila is None:
            time.sleep(1)
    if fila is None:
        log.warning("[8/8] Consulta de Clientes: no hay ninguna fila con Tipo Documento %r.", TIPO_DOC_CLIENTE)
        return None

    log.info("[8/8] Consulta de Clientes: marco la fila -> %s", " | ".join(fila.text.split("\n")))
    radios = fila.find_elements(By.XPATH, ".//input[@type='radio' or @type='checkbox']")
    marca = radios[0] if radios else fila.find_elements(By.TAG_NAME, "td")[0]
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", marca)
    try:
        marca.click()
    except Exception:
        driver.execute_script("arguments[0].click();", marca)   # radio oculto por el estilo
    time.sleep(0.5)
    if radios and not radios[0].is_selected():
        driver.execute_script("arguments[0].checked = true; arguments[0].dispatchEvent(new Event('change', {bubbles: true}));"
                              " arguments[0].dispatchEvent(new MouseEvent('click', {bubbles: true}));", radios[0])

    xp = (_xpath_texto(("button", "a"), "seleccionar") +
          " | //input[contains(translate(@value,'SELECCIONAR','seleccionar'),'seleccionar')]")
    botones = [b for b in driver.find_elements(By.XPATH, xp) if b.is_displayed()]
    if not botones:
        log.warning("[8/8] Consulta de Clientes: no encontré el botón Seleccionar.")
        return None
    antes = driver.window_handles
    log.info("[8/8] Consulta de Clientes: clic en Seleccionar.")
    try:
        botones[0].click()
    except Exception:
        driver.execute_script("arguments[0].click();", botones[0])

    # Consulta de Productos puede cargar en esta misma ventana o abrir otra
    fin = time.time() + TIMEOUT_POPUP
    while time.time() < fin:
        _alerta_nativa(driver)
        nuevas = [h for h in driver.window_handles if h not in antes]
        if nuevas:
            driver.switch_to.window(nuevas[-1])
            log.info("[8/8] Consulta de Productos abierta en una ventana nueva.")
            return nuevas[-1]
        if popup in driver.window_handles:
            try:
                driver.switch_to.window(popup)
                if not REGEX_CONSULTA_CLIENTES.search(driver.execute_script(_JS_TEXTO_VISIBLE) or ""):
                    log.info("[8/8] Consulta de Productos cargó en la misma ventana.")
                    return popup
            except Exception:
                pass                          # está navegando
        time.sleep(1)
    log.warning("[8/8] Consulta de Clientes: tras Seleccionar no apareció Consulta de Productos.")
    return None


_XPATH_ACEPTAR = (_xpath_texto(("button", "a", "span", "div"), "aceptar") +
                  " | //input[contains(translate(@value,'ACEPTR','aceptr'),'aceptar')]")


def aceptar_alerta(driver, handle_alerta, timeout=10):
    """En la ventana de la alerta: busca el botón Aceptar (también en iframes) y le hace clic.
    Si no lo encuentra o la ventana no se cierra, la cierra directamente."""
    driver.switch_to.window(handle_alerta)
    log.info("[8/8] Alerta: controlando la ventana %r | %s", driver.title, driver.current_url)

    boton = _buscar_en_frames(driver, _XPATH_ACEPTAR, timeout=timeout)
    if boton is None:
        visibles = [e.text.strip() for e in driver.find_elements(By.XPATH, "//button | //a | //input")
                    if e.is_displayed() and e.text.strip()][:15]
        log.warning("[8/8] Alerta: no encontré 'Aceptar'. Botones visibles: %s", ", ".join(visibles) or "ninguno")
    else:
        # el xpath puede devolver el span de adentro: subo al botón/enlace si lo hay
        padre = boton.find_elements(By.XPATH, "./ancestor-or-self::*[self::button or self::a][1]")
        boton = padre[0] if padre else boton
        log.info("[8/8] Alerta: botón encontrado (%s %r); clic...", boton.tag_name, boton.text.strip())
        try:
            boton.click()
        except Exception:
            driver.execute_script("arguments[0].click();", boton)
        log.info("[8/8] Alerta: clic en Aceptar. Esperando que se cierre (hasta %s s)...",
                 TIMEOUT_CIERRE_ALERTA)
        if _esperar_cierre_alerta(driver, handle_alerta, TIMEOUT_CIERRE_ALERTA):
            log.info("[8/8] Alerta: cerrada.")
            return

    # no se cerró sola (o no hubo botón): la cierro yo y confirmo que desapareció
    try:
        driver.switch_to.window(handle_alerta)
        driver.close()
        log.info("[8/8] Alerta: no se cerró sola; la cerré desde el navegador.")
    except Exception:
        pass
    if not _esperar_cierre_alerta(driver, handle_alerta, 5):
        log.warning("[8/8] Alerta: la ventana sigue abierta; continúo igual.")


def _esperar_cierre_alerta(driver, handle_alerta, timeout):
    """True cuando la ventana de la alerta desaparece (o deja de mostrar la alerta)."""
    fin = time.time() + timeout
    while time.time() < fin:
        if handle_alerta not in driver.window_handles:
            return True
        try:
            driver.switch_to.window(handle_alerta)
            if not _texto_alerta(driver):         # la misma ventana siguió a otra página
                return True
        except Exception:
            return True                           # se cerró mientras la consultaba
        time.sleep(0.5)
    return False


def extraer_totales(driver, timeout=TIMEOUT_DATOS):
    """Espera a que el popup termine de cargar sus tablas y recién ahí las lee.

    Relee cada pocos segundos: da por buena la lectura cuando dos pasadas seguidas
    devuelven la misma cantidad de filas, así no lee una tabla a medio pintar.
    """
    _esperar_carga(driver)
    log.info("[8/8] Esperando los datos del popup (hasta %s s)...", timeout)

    fin = time.time() + timeout
    filas, previas = [], -1
    while time.time() < fin:
        filas = _leer_tablas(driver)
        if filas and len(filas) == previas:
            log.info("[8/8] Datos completos: %d filas de servicios.", len(filas))
            return filas
        if filas:
            log.info("[8/8] Cargando... %d filas hasta ahora.", len(filas))
        previas = len(filas)
        time.sleep(ESPERA_DATOS)

    if filas:
        log.warning("[8/8] Se agotó el tiempo; me quedo con las %d filas leídas.", len(filas))
    return filas


def inspeccionar_popup(driver, max_filas=3):
    """Registra en el log qué trae el popup: título, columnas y primeras filas."""
    log.info("[8/8] Detalle -> %r | %s", driver.title, driver.current_url)
    for i, tabla in enumerate(driver.find_elements(By.TAG_NAME, "table")):
        try:
            cabecera = [" ".join(th.text.split()) for th in tabla.find_elements(By.TAG_NAME, "th")]
            filas = tabla.find_elements(By.XPATH, ".//tbody/tr")
        except Exception:
            continue
        if not cabecera and not filas:
            continue
        log.info("[8/8]   tabla %d (%d filas): %s", i, len(filas), " | ".join(cabecera) or "(sin encabezados)")
        for fila in filas[:max_filas]:
            celdas = [_valor_celda(c) for c in fila.find_elements(By.TAG_NAME, "td")]
            if any(celdas):
                log.info("[8/8]     %s", " | ".join(celdas))


def _contenedor_detalle(driver):
    """Si el detalle se abrió como ventana modal dentro de la página, esa modal; si no, la página.
    Así no se clica un número de la tabla de productos que queda debajo."""
    for modal in reversed(driver.find_elements(
            By.CSS_SELECTOR, ".modal.in, .modal.show, .modal[style*='display: block'], [role='dialog']")):
        try:
            if modal.is_displayed() and REGEX_DETALLE.search(modal.text or ""):
                return modal
        except Exception:
            continue
    return driver


def _primer_numero(driver):
    """El primer número clicable de la ventana de detalle (enlace, botón o celda), o None."""
    xp = (".//tbody//tr//a | .//tbody//tr//button | .//tbody//tr//input"
          " | .//tbody//tr//span | .//tbody//tr//td")
    base = _contenedor_detalle(driver)
    if base is driver:
        xp = xp.replace(".//", "//")
    for el in base.find_elements(By.XPATH, xp):
        try:
            texto = (el.text or el.get_attribute("value") or "").strip()
            if re.fullmatch(r"\d[\d\s.\-]*", texto) and el.is_displayed():
                return el, texto
        except Exception:
            continue                          # se repintó mientras lo leía
    return None, ""


def clic_primer_numero(driver, handle_siac, timeout=TIMEOUT_DATOS):
    """En la ventana de detalle: clic en el primer número, que redirige a SIAC; espera ahí."""
    log.info("[8/8] Detalle: esperando el primer número (hasta %s s)...", timeout)
    fin = time.time() + timeout
    el, texto = None, ""
    while time.time() < fin and el is None:
        el, texto = _primer_numero(driver)
        if el is None:
            time.sleep(1)
    if el is None:
        log.warning("[8/8] Detalle: no encontré ningún número que clicar.")
        return False

    log.info("[8/8] Detalle: clic en el primer número -> %s", texto)
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
    time.sleep(0.5)
    try:
        el.click()
    except Exception:
        driver.execute_script("arguments[0].click();", el)

    driver.switch_to.window(handle_siac)
    log.info("[8/8] De vuelta en SIAC; espero %s s...", ESPERA_TRAS_NUMERO)
    time.sleep(ESPERA_TRAS_NUMERO)
    _esperar_carga(driver)
    log.info("[8/8] SIAC: %s", driver.current_url)
    return True


def abrir_detalle_activas(driver, handle_popup):
    """Clic en el Total Activas de la 1ra fila; si no es >= 1, en el recuadro de su derecha
    (Total No Activas). Si la fila tiene 0 y 0, prueba con la fila siguiente.
    Devuelve el handle del popup de líneas que se abre, o None."""
    driver.switch_to.window(handle_popup)
    totales = esperar_primera_fila(driver)
    if totales is None:
        return None

    eleccion = None
    for n, (activas, no_activas) in enumerate(totales):
        if activas >= 1:
            eleccion = (n, "activas", "Total Activas", activas)
        elif no_activas >= 1:
            eleccion = (n, "no_activas", "Total No Activas", no_activas)
            log.info("[8/8] Fila %d: Total Activas = %d; uso el de la derecha, Total No Activas = %d.",
                     n + 1, activas, no_activas)
        if eleccion:
            break
        log.info("[8/8] Fila %d: 0 activas y 0 no activas; paso a la siguiente.", n + 1)
    if eleccion is None:
        log.info("[8/8] Ninguna fila tiene activas ni no activas; no hay detalle que abrir.")
        return None
    n, clave, columna, valor = eleccion

    handles_antes = driver.window_handles
    log.info("[8/8] Esperando el recuadro %s = %d de la fila %d (hasta %s s)...",
             columna, valor, n + 1, TIMEOUT_BOTON_TOTAL)
    detalle = clic_total_y_esperar_popup(driver, clave, handles_antes, n)
    if detalle is None:
        log.warning("[8/8] Ninguna forma de clic en %s abrió el popup.", columna)
        try:
            driver.switch_to.window(handle_popup)
        except Exception:
            pass
        return None

    driver.switch_to.window(detalle)
    _esperar_carga(driver)
    time.sleep(ESPERA_DATOS)
    inspeccionar_popup(driver)
    return detalle


def cerrar_popup(driver, handle_popup, handle_siac):
    """Clic en Cerrar; si no aparece el botón, cierra la ventana directamente."""
    if handle_popup not in driver.window_handles:  # ya se cerró sola
        try:
            driver.switch_to.window(handle_siac)
        except Exception:
            pass
        return
    try:
        # el botón Cerrar se busca SOLO en el popup (en SIAC podría ser "Cerrar sesión")
        driver.switch_to.window(handle_popup)
        xp = (_xpath_texto(("button", "a", "span"), "cerrar") +
              " | //input[contains(translate(@value,'CERRAR','cerrar'),'cerrar')]")
        botones = [e for e in driver.find_elements(By.XPATH, xp) if e.is_displayed()]
        if botones:
            log.info("[8/8] Cerrando el popup con el botón Cerrar.")
            try:
                botones[0].click()
            except Exception:
                driver.execute_script("arguments[0].click();", botones[0])
            time.sleep(1.5)
    except Exception as e:
        log.debug("Botón Cerrar no utilizable: %s", e)

    if handle_popup in driver.window_handles:      # el botón no la cerró
        try:
            driver.switch_to.window(handle_popup)
            driver.close()
            log.info("[8/8] Popup cerrado por el navegador.")
        except Exception:
            pass

    try:
        driver.switch_to.window(handle_siac)
    except Exception:
        vivas = _ventanas_vivas(driver)
        if vivas:
            driver.switch_to.window(vivas[-1][0])


# Lee el panel "Datos del Cliente": busca la etiqueta (texto exacto, sin tildes ni
# mayúsculas) y devuelve el texto del elemento de al lado.
_JS_DATOS_CLIENTE = r"""
const norm = s => (s || '').replace(/\s+/g, ' ').trim().toLowerCase()
                           .normalize('NFD').replace(/[̀-ͯ]/g, '');
const els = document.querySelectorAll('td,th,label,dt,dd,div,span,strong,b,p');
function valor(etiquetas) {
  for (const et of etiquetas) {
    for (const el of els) {
      if (!el.offsetParent || norm(el.innerText) !== et) continue;
      const sib = el.nextElementSibling || (el.parentElement && el.parentElement.nextElementSibling);
      const v = sib ? sib.innerText.replace(/\s+/g, ' ').trim() : '';
      if (v) return v;
    }
  }
  return '';
}
return {
  cliente:             valor(['cliente', 'razon social', 'nombres y apellidos']),
  contacto:            valor(['contacto']),
  representante_legal: valor(['representante legal', 'rep. legal']),
  telefono:            valor(['tel. referencia 1', 'tel. referencia', 'telefono referencia', 'telefono']),
  email:               valor(['email', 'e-mail', 'correo', 'correo electronico']),
};
"""

# Datos del cliente que se guardan, en el orden de las columnas del CSV
CAMPOS_CLIENTE = ["cliente", "contacto", "representante_legal", "telefono", "email"]


def leer_datos_cliente(driver, timeout=TIMEOUT_DATOS):
    """En la ficha del cliente de SIAC: un dict con CAMPOS_CLIENTE (vacíos si no están)."""
    fin = time.time() + timeout
    datos = {}
    while time.time() < fin:
        try:
            datos = driver.execute_script(_JS_DATOS_CLIENTE) or {}
        except Exception:
            pass                              # la página aún se está pintando
        if datos.get("cliente"):
            break
        time.sleep(1)
    datos = {c: datos.get(c, "") for c in CAMPOS_CLIENTE}
    datos["cliente"] = re.sub(r"\s*\.\.\.\s*$", "", datos["cliente"])   # quita el botón "..."
    log.info("[8/8] %s", " | ".join(f"{c}: {datos[c]!r}" for c in CAMPOS_CLIENTE))
    return datos


# Estado de cada RUC en el CSV (columna "estado"); "detalle" dice el motivo concreto
OK = "OK"                  # se leyó la ficha del cliente
SIN_DATOS = "SIN_DATOS"    # SIAC no tiene datos: "No existen datos", sin opción RUC, 0 líneas...
BLOQUEO = "BLOQUEO"        # alerta "No está autorizado"
FALLO = "FALLO"            # problema del robot: no abrió el popup/detalle, tiempo agotado, excepción
ESTADOS = [OK, SIN_DATOS, BLOQUEO, FALLO]

# Columna "estado_ruc": la versión simple del estado
VALIDO, VACIO, BLOQUEADO = "VALIDO", "VACIO", "BLOQUEADO"
ESTADO_RUC = {OK: VALIDO, SIN_DATOS: VACIO, FALLO: VACIO, BLOQUEO: BLOQUEADO}

COLUMNAS_CSV = ["ruc", "estado_ruc"] + CAMPOS_CLIENTE + ["estado", "detalle"]


def _leer_fila(cabecera, fila):
    """(ruc, campos, estado, detalle) de una fila, en cualquier formato del CSV (lee por nombre
    de columna). Si la fila es de un formato sin estado, lo deduce: OK si tiene cliente,
    FALLO si dice ERROR, SIN_DATOS si está vacía."""
    d = dict(zip(cabecera, fila))
    campos = {c: d.get(c, "") for c in CAMPOS_CLIENTE}
    estado, detalle = d.get("estado", ""), d.get("detalle", "")
    if not estado:
        if campos["cliente"] == "ERROR":
            campos = {c: "" for c in CAMPOS_CLIENTE}
            estado, detalle = FALLO, "ERROR (formato anterior)"
        elif campos["cliente"]:
            estado = OK
        else:
            estado, detalle = SIN_DATOS, "vacío, motivo no registrado (formato anterior)"
    return d.get("ruc", "").strip(), campos, estado, detalle


def _fila_csv(ruc, campos, estado, detalle):
    return ([ruc, ESTADO_RUC.get(estado, VACIO)] + [campos.get(c, "") for c in CAMPOS_CLIENTE]
            + [estado, detalle])


def guardar_fila_csv(ruc, datos=None, estado=OK, detalle="", ruta=ARCHIVO_CSV):
    """Agrega una fila al CSV: ruc + estado_ruc + CAMPOS_CLIENTE (de 'datos') + estado + detalle."""
    archivo = Path(ruta)
    if archivo.exists():
        with open(archivo, newline="", encoding="utf-8-sig") as f:
            filas = list(csv.reader(f))
        cabecera = filas[0] if filas else []
        if cabecera != COLUMNAS_CSV and "ruc" in cabecera and set(CAMPOS_CLIENTE) <= set(cabecera):
            # formato anterior: lo paso al nuevo sin perder el avance (ultimo_ruc_csv sigue funcionando)
            with open(archivo, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                w.writerow(COLUMNAS_CSV)
                for fila in filas[1:]:
                    r, campos, est, det = _leer_fila(cabecera, fila)
                    if r:
                        w.writerow(_fila_csv(r, campos, est, det))
            log.info("[8/8] Pasé el CSV existente al formato nuevo (%s).", ", ".join(COLUMNAS_CSV))
        elif cabecera != COLUMNAS_CSV:        # CSV de otro formato: lo aparto, no lo mezclo
            respaldo = archivo.with_name(f"{archivo.stem}_anterior_{time.strftime('%Y%m%d_%H%M%S')}.csv")
            archivo.rename(respaldo)
            log.info("[8/8] El CSV tenía otras columnas; lo moví a %s", respaldo.name)
    nuevo = not archivo.exists()
    with open(archivo, "a", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        if nuevo:
            w.writerow(COLUMNAS_CSV)
        w.writerow(_fila_csv(ruc, datos or {}, estado, detalle))


def metricas_csv(ruta=ARCHIVO_CSV, max_listado=20):
    """Resumen del CSV: cuántos salieron válidos, vacíos (por falta de datos o fallo del robot)
    y bloqueados. Si un RUC se repite, cuenta su última fila."""
    archivo = Path(ruta)
    if not archivo.exists():
        log.info("[métricas] No existe %s todavía.", archivo.resolve())
        return {}
    with open(archivo, newline="", encoding="utf-8-sig") as f:
        filas = list(csv.reader(f))
    if not filas:
        return {}
    cabecera, filas = filas[0], filas[1:]

    ultimas = {}                              # ruc -> (campos, estado, detalle)
    for fila in filas:
        ruc, campos, est, det = _leer_fila(cabecera, fila)
        if ruc:
            ultimas[ruc] = (campos, est, det)

    total = len(ultimas)
    por_estado = {e: [] for e in ESTADOS}
    for ruc, (_, est, det) in ultimas.items():
        por_estado.setdefault(est, []).append((ruc, det))
    n = {e: len(v) for e, v in por_estado.items()}
    pct = lambda k: f"{100 * k / total:5.1f}%" if total else "  -  "
    vacios = sum(v for e, v in n.items() if ESTADO_RUC.get(e, VACIO) == VACIO)

    log.info("[métricas] ===== Resumen de %s =====", archivo.name)
    log.info("[métricas] RUCs procesados: %d (filas en el CSV: %d)", total, len(filas))
    log.info("[métricas]   %-9s ............ %5d  %s", VALIDO, n[OK], pct(n[OK]))
    log.info("[métricas]   %-9s ............ %5d  %s", VACIO, vacios, pct(vacios))
    log.info("[métricas]     - falta de datos .... %5d  %s", n[SIN_DATOS], pct(n[SIN_DATOS]))
    log.info("[métricas]     - fallo del robot ... %5d  %s", n[FALLO], pct(n[FALLO]))
    for e in (e for e in n if e not in ESTADOS and n[e]):
        log.info("[métricas]     - %s ... %5d  %s", e, n[e], pct(n[e]))
    log.info("[métricas]   %-9s ............ %5d  %s", BLOQUEADO, n[BLOQUEO], pct(n[BLOQUEO]))

    # de los válidos, cuántos vinieron con campos en blanco
    if por_estado[OK]:
        incompletos = {c: sum(1 for ruc, _ in por_estado[OK] if not ultimas[ruc][0].get(c))
                       for c in CAMPOS_CLIENTE if c != "cliente"}
        log.info("[métricas]   Válidos sin: %s",
                 ", ".join(f"{c} {k}" for c, k in incompletos.items()))

    # el motivo concreto dentro de cada grupo, y qué RUCs son
    for est, titulo in ((SIN_DATOS, "Vacíos por falta de datos"), (FALLO, "Vacíos por fallo del robot"),
                        (BLOQUEO, "Bloqueados")):
        lista = por_estado[est]
        if not lista:
            continue
        motivos = {}
        for _, det in lista:
            motivos[det or "(sin detalle)"] = motivos.get(det or "(sin detalle)", 0) + 1
        log.info("[métricas] %s (%d):", titulo, len(lista))
        for det, k in sorted(motivos.items(), key=lambda x: -x[1]):
            log.info("[métricas]     %4d  %s", k, det[:100])
        rucs = [ruc for ruc, _ in lista]
        resto = f" ... y {len(rucs) - max_listado} más" if len(rucs) > max_listado else ""
        log.info("[métricas]   RUCs: %s%s", ", ".join(rucs[:max_listado]), resto)
    if total - n[OK]:
        log.info("[métricas] (filtra la columna 'estado_ruc' o 'estado' del CSV para ver todos los RUCs)")

    return {e: [ruc for ruc, _ in v] for e, v in por_estado.items()}


def ultimo_ruc_csv(ruta=ARCHIVO_CSV):
    """El RUC de la última fila del CSV (donde se quedó la ejecución anterior), o None."""
    archivo = Path(ruta)
    if not archivo.exists():
        return None
    ultimo = None
    with open(archivo, newline="", encoding="utf-8-sig") as f:
        filas = csv.reader(f)
        next(filas, None)                     # cabecera
        for fila in filas:
            if fila and fila[0].strip():
                ultimo = fila[0].strip()
    return ultimo


_caidas_por_ruc = {}                          # ruc -> veces que se cerró el navegador con él


def _caida_con_ruc(ruc, error):
    """El navegador se cerró buscando 'ruc'. Si ya pasó MAX_CAIDAS_POR_RUC veces, lo guardo como
    FALLO para no quedar atascado en él. Siempre lanza NavegadorCerrado."""
    _caidas_por_ruc[ruc] = _caidas_por_ruc.get(ruc, 0) + 1
    veces = _caidas_por_ruc[ruc]
    if veces >= MAX_CAIDAS_POR_RUC:
        log.warning("[8/8] %s: el navegador se cerró %d veces con este RUC; lo guardo como %s.",
                    ruc, veces, FALLO)
        guardar_fila_csv(ruc, None, FALLO, f"el navegador se cerró {veces} veces con este RUC")
    raise NavegadorCerrado(f"se cerró el navegador buscando {ruc} ({type(error).__name__}: {error})"[:300])


def buscar_lote(driver, handle_portal, handle_siac, valores=None):
    """Busca en SIAC cada valor del Excel, uno tras otro.

    Retoma donde se quedó: si el CSV ya tiene filas, empieza en el RUC que sigue al último guardado.
    Tras guardar cada valor en el CSV cierra todas las ventanas (incluida la de SIAC Único),
    vuelve al portal y hace clic otra vez en SIAC Único para el siguiente valor."""
    valores = leer_valores() if valores is None else valores
    if not valores:
        log.warning("[8/8] El Excel no tiene valores que buscar.")
        return

    total = len(valores)
    inicio = 0
    ultimo = ultimo_ruc_csv()
    if ultimo in valores:
        inicio = valores.index(ultimo) + 1
        log.info("[8/8] El CSV terminó en %s (%d/%d); retomo desde el siguiente.", ultimo, inicio, total)
    elif ultimo:
        log.warning("[8/8] El último RUC del CSV (%s) no está en el Excel; empiezo desde el primero.", ultimo)
    if inicio >= total:
        log.info("[8/8] Todos los valores del Excel ya están en el CSV; no hay nada que buscar.")
        metricas_csv()
        return

    log.info("[8/8] Lote de %d búsquedas (quedan %d). Resultados -> %s",
             total, total - inicio, Path(ARCHIVO_CSV).resolve())

    for i, valor in enumerate(valores[inicio:], inicio + 1):
        log.info("[8/8] (%d/%d) --- %s ---", i, total, valor)
        try:
            handles_antes = driver.window_handles
            buscar_en_siac(driver, valor)

            datos = {}
            estado, detalle = FALLO, ""
            popup = esperar_popup(driver, handles_antes)
            if popup is None:
                log.warning("[8/8] (%d/%d) %s: no apareció el popup en %s s.", i, total, valor, TIMEOUT_POPUP)
                detalle = f"no apareció el popup en {TIMEOUT_POPUP} s"
            else:
                driver.switch_to.window(popup)
                _esperar_carga(driver)
                contenido, mensaje, h_alerta = esperar_contenido_popup(driver, popup,
                                                                       (handle_siac, handle_portal))
                if contenido == "clientes":
                    # varios clientes con ese número: elijo el de RUC y sigo con sus productos
                    elegido = elegir_cliente_ruc(driver, popup)
                    if elegido is None:
                        contenido = "sin_ruc"
                    else:
                        popup = elegido
                        _esperar_carga(driver)
                        contenido, mensaje, h_alerta = esperar_contenido_popup(driver, popup,
                                                                               (handle_siac, handle_portal))
                if contenido == "datos":
                    driver.switch_to.window(popup)
                    filas = extraer_totales(driver)
                else:
                    filas = []

                if contenido in ("alerta", "nativa"):
                    log.warning("[8/8] (%d/%d) %s: alerta -> %r. Acepto y guardo vacío.",
                                i, total, valor, mensaje)
                    if h_alerta:
                        aceptar_alerta(driver, h_alerta)
                    # cierro cualquier otra ventana que haya quedado (ej. Consulta de Productos vacía)
                    for h in list(driver.window_handles):
                        if h not in (handle_siac, handle_portal):
                            cerrar_popup(driver, h, handle_siac)
                    driver.switch_to.window(handle_siac)      # la alerta ya cerró: vuelvo a SIAC
                    _esperar_carga(driver)
                    estado, detalle = BLOQUEO, mensaje or "alerta"
                elif contenido == "sin_ruc":
                    log.info("[8/8] (%d/%d) %s: Consulta de Clientes sin opción %r. Guardo vacío.",
                             i, total, valor, TIPO_DOC_CLIENTE)
                    estado, detalle = SIN_DATOS, f"Consulta de Clientes sin opción {TIPO_DOC_CLIENTE}"
                elif contenido == "vacio":
                    log.info("[8/8] (%d/%d) %s: popup sin productos. Cierro, guardo vacío y paso al siguiente.",
                             i, total, valor)
                    estado, detalle = SIN_DATOS, "popup sin productos"
                elif contenido is None:
                    log.info("[8/8] (%d/%d) %s: el popup no mostró datos ni alerta a tiempo.", i, total, valor)
                    detalle = f"el popup no cargó en {TIMEOUT_DATOS} s"
                elif not filas:
                    log.info("[8/8] (%d/%d) %s: el popup vino vacío.", i, total, valor)
                    estado, detalle = SIN_DATOS, "popup sin filas de servicios"
                elif all(act == 0 and no == 0 for _, act, no in filas):
                    log.info("[8/8] (%d/%d) %s: todas las filas tienen 0 activas y 0 no activas.",
                             i, total, valor)
                    estado, detalle = SIN_DATOS, "0 líneas activas y 0 no activas"
                elif CLIC_EN_TOTAL_ACTIVAS:
                    h_detalle = abrir_detalle_activas(driver, popup)
                    if h_detalle is None:
                        detalle = "no se abrió el detalle de líneas"
                    elif not clic_primer_numero(driver, handle_siac):   # redirige a SIAC
                        detalle = "el detalle no tenía número que clicar"
                    else:
                        datos = leer_datos_cliente(driver)
                        cerrar_popup(driver, h_detalle, handle_siac)   # cierro el detalle si sigue abierto
                        if datos.get("cliente"):
                            estado, detalle = OK, ""
                        else:
                            detalle = "la ficha del cliente no mostró el nombre"
                else:
                    detalle = "CLIC_EN_TOTAL_ACTIVAS desactivado: no se lee la ficha"
                cerrar_popup(driver, popup, handle_siac)

            if estado != OK and not _driver_vivo(driver):
                # el "no encontré nada" fue porque se cerró el navegador: no lo guardo como vacío
                raise NavegadorCerrado("el navegador se cerró durante la búsqueda")
            guardar_fila_csv(valor, datos, estado, detalle)
            log.info("[8/8] (%d/%d) %s guardado en el CSV: %s%s", i, total, valor, estado,
                     f" ({detalle})" if detalle else "")
        except Exception as e:
            if not _driver_vivo(driver):
                _caida_con_ruc(valor, e)      # lanza NavegadorCerrado: se reabre y se reintenta este RUC
            log.warning("[8/8] (%d/%d) %s falló: %s. Lo guardo como %s y sigo.", i, total, valor, e, FALLO)
            guardar_fila_csv(valor, None, FALLO, f"excepción: {type(e).__name__}: {e}"[:200])

        if i < total:
            # cierro todos los popups y SIAC Único, y lo abro de nuevo desde el portal
            log.info("[8/8] Preparando el siguiente RUC: cierro todo y reabro %r...", TEXTO_OPCION)
            try:
                volver_al_portal(driver, handle_portal)
                time.sleep(PAUSA_ENTRE_BUSQUEDAS)
                handle_siac = abrir_siac(driver)
            except Exception as e2:
                # el RUC ya quedó guardado: al reabrir el navegador se sigue con el siguiente
                raise NavegadorCerrado(f"no pude reabrir SIAC: {e2}") from e2

    log.info("[8/8] Lote terminado (%d valores). CSV: %s", total, Path(ARCHIVO_CSV).resolve())
    metricas_csv()


def _ventanas_vivas(driver):
    """(handle, título, url) de las ventanas que todavía responden."""
    vivas = []
    for h in driver.window_handles:
        try:
            driver.switch_to.window(h)
            vivas.append((h, driver.title or "", driver.current_url or ""))
        except Exception:
            continue          # se cerró mientras la consultaba
    return vivas


def ir_a_ventana_siac(driver, pestanas_antes, timeout=TIMEOUT_VENTANA_SIAC):
    """Se cambia a la ventana de SIAC, tolerando popups que se abren y se cierran."""
    fin = time.time() + timeout
    nueva = None
    while time.time() < fin:
        vivas = _ventanas_vivas(driver)
        for h, titulo, url in vivas:
            if REGEX_VENTANA_SIAC.search(url) or REGEX_VENTANA_SIAC.search(titulo):
                driver.switch_to.window(h)
                log.info("[7/8] Ventana de SIAC: %r -> %s", titulo, url)
                return True
        nuevas = [v for v in vivas if v[0] not in pestanas_antes]
        if nuevas:
            nueva = nuevas[-1]     # existe pero aún no resuelve su URL
        time.sleep(1)

    vivas = _ventanas_vivas(driver)
    if not vivas:
        raise RuntimeError("Se cerraron todas las ventanas del navegador.")

    if nueva and any(v[0] == nueva[0] for v in vivas):
        driver.switch_to.window(nueva[0])
        log.warning("[7/8] Uso la ventana nueva %r; no confirmé que sea SIAC.", nueva[1])
        return False

    log.warning("[7/8] No apareció la ventana de SIAC en %s s. Ventanas vivas:", timeout)
    for _, titulo, url in vivas:
        log.warning("[7/8]   %r -> %s", titulo, url)
    driver.switch_to.window(vivas[-1][0])
    return False


def _esperar_carga(driver, timeout=TIMEOUT_WEB):
    """readyState == complete. False si la ventana murió (no lanza excepción)."""
    fin = time.time() + timeout
    while time.time() < fin:
        try:
            if driver.execute_script("return document.readyState") == "complete":
                return True
        except Exception:
            return False
        time.sleep(0.5)
    return False


def clic_opcion_portal(driver, texto=TEXTO_OPCION):
    """Desde el portal abre SIAC Único y lanza la búsqueda (un valor o el lote del Excel)."""
    handle_portal = driver.current_window_handle
    handle_siac = abrir_siac(driver, texto)
    if VALOR_BUSQUEDA:
        buscar_en_siac(driver, VALOR_BUSQUEDA)     # override: un solo valor
    else:
        buscar_lote(driver, handle_portal, handle_siac)   # itera el Excel


def volver_al_portal(driver, handle_portal):
    """Cierra TODAS las ventanas (popups, detalle y la de SIAC Único) menos la del portal."""
    _alerta_nativa(driver)
    cerradas = 0
    for h in list(driver.window_handles):
        if h == handle_portal:
            continue
        try:
            driver.switch_to.window(h)
            _alerta_nativa(driver)
            driver.close()
            cerradas += 1
        except Exception:
            continue                          # ya se había cerrado
    if handle_portal not in driver.window_handles:
        raise RuntimeError("Se cerró la ventana del portal.")
    driver.switch_to.window(handle_portal)
    driver.switch_to.default_content()
    if REGEX_VENTANA_SIAC.search(driver.current_url or ""):
        # SIAC se abrió en la misma pestaña del portal: vuelvo a cargar el portal
        driver.get(URL_SELENIUM)
        esperar_credenciales_usuario(driver)
        _esperar_carga(driver)
    log.info("[7/8] Cerré %d ventana(s); de vuelta en el portal.", cerradas)


def abrir_siac(driver, texto=TEXTO_OPCION):
    """Busca la opción (enlace, botón o texto) dentro del portal y le hace clic.
    Espera la ventana de SIAC, deja el desplegable listo y devuelve su handle."""
    # Normaliza a minúsculas y sin tildes para comparar (el portal escribe "Unico" y "Único")
    minus = "ABCDEFGHIJKLMNOPQRSTUVWXYZÁÉÍÓÚÜÑ", "abcdefghijklmnopqrstuvwxyzaeiouun"
    objetivo = texto.lower().translate(str.maketrans("áéíóúüñ", "aeiouun"))
    tr = f"translate(normalize-space(.),'{minus[0]}','{minus[1]}')"
    tr_attr = lambda a: f"translate(normalize-space(@{a}),'{minus[0]}','{minus[1]}')"
    # Solo el elemento más interno con el texto: un <td>/<li> contenedor del menú también
    # "contiene" el texto, y al clicar su centro caía en otro enlace (ej. SIVCO).
    hoja = f"contains({tr},'{objetivo}') and not(.//*[contains({tr},'{objetivo}')])"
    xpath = (
        f"//a[{hoja}]"
        f" | //button[{hoja}]"
        f" | //span[{hoja}]"
        f" | //td[{hoja}]"
        f" | //li[{hoja}]"
        f" | //*[contains({tr_attr('title')},'{objetivo}') or contains({tr_attr('alt')},'{objetivo}')]"
    )
    el = _buscar_en_frames(driver, xpath)

    if el is None:
        enlaces = sorted({a.text.strip() for a in driver.find_elements(By.TAG_NAME, "a")
                          if a.text.strip() and len(a.text.strip()) < 60})[:60]
        raise RuntimeError(
            f"No encontré la opción '{texto}' en el portal. "
            f"Enlaces visibles: {', '.join(enlaces) or 'ninguno (¿el portal cargó?)'}"
        )

    log.info("[7/8] Portal: clic en <%s> %r (href=%s)", el.tag_name,
             el.text.strip() or texto, el.get_attribute("href"))
    pestanas_antes = driver.window_handles
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
    time.sleep(0.3)
    try:
        el.click()
    except Exception:
        driver.execute_script("arguments[0].click();", el)  # por si algo la tapa

    ir_a_ventana_siac(driver, pestanas_antes)

    # SIAC pide autenticación otra vez (es otra aplicación tras el portal)
    time.sleep(2)
    esperar_credenciales_usuario(driver)

    if not _esperar_carga(driver):
        # El popup intermedio se cerró al redirigir: busco la ventana definitiva
        log.warning("[7/8] La ventana se cerró mientras cargaba; busco la de SIAC otra vez.")
        ir_a_ventana_siac(driver, pestanas_antes)
        esperar_credenciales_usuario(driver)
        _esperar_carga(driver)
    log.info("[7/8] SIAC listo: %s", driver.current_url)

    time.sleep(2)                       # que SIAC termine de pintar la barra de búsqueda
    seleccionar_tipo_busqueda(driver)
    return driver.current_window_handle


def _credenciales_portal():
    if REUSAR_CREDENCIALES_VPN:
        pwd = keyring.get_password(KEYRING_SERVICIO, USUARIO_VPN)   # la misma de la VPN
        origen = "--guardar-password"
    else:
        pwd = keyring.get_password(KEYRING_SERVICIO_PORTAL, USUARIO_PORTAL)
        origen = "--guardar-password-portal"
    if not pwd:
        raise RuntimeError(f"No hay contraseña guardada. Ejecuta: python toques_y_selenium.py {origen}")
    return USUARIO_PORTAL, pwd


def auth_basica(driver):
    """Manda las credenciales en la cabecera Authorization para evitar el popup del 401."""
    usuario, pwd = _credenciales_portal()
    token = base64.b64encode(f"{usuario}:{pwd}".encode()).decode()
    del pwd
    driver.execute_cdp_cmd("Network.enable", {})
    driver.execute_cdp_cmd("Network.setExtraHTTPHeaders", {"headers": {"Authorization": f"Basic {token}"}})
    log.info("Cabecera de autenticación configurada para el portal.")


def _ventana_navegador():
    for w in _escritorio().windows():
        try:
            marca = "Edge" if NAVEGADOR == "edge" else "Chrome"
            if w.class_name() == "Chrome_WidgetWin_1" and w.is_visible() and marca in (w.window_text() or ""):
                return w
        except Exception:
            continue
    return None


def _hay_dialogo_auth(driver):
    """El popup bloquea la carga: la página queda sin contenido."""
    try:
        return len(driver.find_elements(By.CSS_SELECTOR, "body *")) == 0
    except Exception:
        return True


def llenar_credenciales_lento(driver):
    """Escribe usuario y contraseña en el popup, tecla por tecla y con pausas largas."""
    ventana = _ventana_navegador()
    if ventana is None:
        log.warning("[6/8] No encontré la ventana del navegador.")
        return False

    usuario, pwd = _credenciales_portal()
    log.info("[6/8] Portal: tecleando credenciales de %s (%.2fs por tecla)...", usuario, PAUSA_TECLA)

    ventana.set_focus()
    time.sleep(ESPERA_FOCO)          # el diálogo tarda en tomar el foco

    for ch in usuario:               # tecla por tecla, no de golpe
        send_keys(_escapar_teclas(ch))
        time.sleep(PAUSA_TECLA)

    time.sleep(1)
    send_keys("{TAB}")
    time.sleep(1)

    for ch in pwd:
        send_keys(_escapar_teclas(ch))
        time.sleep(PAUSA_TECLA)
    del pwd

    time.sleep(1)
    send_keys("{ENTER}")
    log.info("[6/8] Portal: credenciales enviadas, esperando respuesta...")

    for _ in range(15):              # hasta 15s a que el portal responda
        if not _hay_dialogo_auth(driver):
            log.info("[6/8] Portal: credenciales ACEPTADAS.")
            return True
        time.sleep(1)
    return False


def esperar_credenciales_usuario(driver):
    """Si sale el popup de autenticación, lo llena automáticamente."""
    if not _hay_dialogo_auth(driver):
        log.info("[6/8] Portal: cargó sin pedir credenciales.")
        return True

    if llenar_credenciales_lento(driver):
        return True

    log.warning("[6/8] El login automático no funcionó; continúo igual.")
    return False


def verificar_password():
    """Diagnostica el esquema de autenticación del portal y prueba la contraseña guardada."""
    guardada = keyring.get_password(KEYRING_SERVICIO, USUARIO_VPN)
    if not guardada:
        print(f"No hay contraseña guardada para {USUARIO_VPN}.")
        print("Ejecuta: python toques_y_selenium.py --guardar-password")
        return

    print(f"Guardada para {USUARIO_VPN}: {guardada[0]}{'*' * (len(guardada) - 2)}{guardada[-1]}"
          f"  ({len(guardada)} caracteres)")

    try:
        import requests
        from requests.auth import HTTPBasicAuth
    except ImportError:
        print("Instala requests: pip install requests")
        return

    # 1) ¿Qué esquema pide el servidor?
    try:
        r = requests.get(URL_SELENIUM, timeout=20)
    except Exception as e:
        print(f"No se pudo conectar ({e}). ¿Está la VPN arriba?")
        return

    esquemas = r.headers.get("WWW-Authenticate", "")
    print(f"El portal responde {r.status_code}; esquemas que pide: {esquemas or '(ninguno)'}")

    usa_ntlm = "ntlm" in esquemas.lower() or "negotiate" in esquemas.lower()
    if usa_ntlm:
        print("=> Usa NTLM/Negotiate, no Basic. Por eso la cabecera Basic no funciona.")

    # 2) Probar las credenciales con el esquema correcto
    intentos = []
    if usa_ntlm:
        try:
            from requests_ntlm import HttpNtlmAuth
        except ImportError:
            print("Para probar NTLM: pip install requests-ntlm")
            return
        for usr in (USUARIO_PORTAL, f"TIM\\{USUARIO_VPN}", f"{USUARIO_VPN}@tim.com.pe"):
            intentos.append((usr, HttpNtlmAuth(usr, guardada)))
    else:
        intentos.append((USUARIO_PORTAL, HTTPBasicAuth(USUARIO_PORTAL, guardada)))

    for usr, auth in intentos:
        try:
            r = requests.get(URL_SELENIUM, auth=auth, timeout=30)
        except Exception as e:
            print(f"  {usr}: error ({e})")
            continue
        if r.status_code < 400:
            print(f"  {usr}: ACEPTADO ({r.status_code}). Usa este valor en USUARIO_PORTAL.")
            return
        print(f"  {usr}: rechazado ({r.status_code})")

    print("Ninguna variante funcionó. Puede ser un dominio distinto al que asumí ('TIM').")


def guardar_password_portal():
    import getpass
    pwd = getpass.getpass(f"Contraseña del portal para {USUARIO_PORTAL} (no se mostrará): ")
    if not pwd:
        print("Contraseña vacía; no se guardó nada.")
        return
    keyring.set_password(KEYRING_SERVICIO_PORTAL, USUARIO_PORTAL, pwd)
    print("Guardada en el Administrador de credenciales de Windows.")


def flujo_portal(url=URL_SELENIUM, opcion=TEXTO_OPCION):
    """Solo la parte de Selenium: abre el portal, clica la opción y devuelve el driver.

    Uso independiente (con la VPN ya conectada):
        from toques_y_selenium import flujo_portal
        driver = flujo_portal()
    """
    driver = abrir_selenium(url, navegar=False)
    driver.get(url)                        # con pageLoadStrategy 'none' vuelve enseguida
    esperar_credenciales_usuario(driver)
    _esperar_carga(driver)
    log.info("[6/8] Portal listo: %s", driver.current_url)
    clic_opcion_portal(driver, opcion)
    return driver


def _descartar_navegador():
    """Cierra lo que quede del navegador caído (libera el perfil para abrir otro)."""
    global _driver_activo
    if _driver_activo is not None:
        try:
            _driver_activo.quit()
        except Exception:
            pass                              # ya estaba muerto
    _driver_activo = None
    time.sleep(3)                             # que Chrome suelte el perfil (--user-data-dir)


def flujo_con_reinicios(url=URL_SELENIUM, opcion=TEXTO_OPCION):
    """flujo_portal(), pero si el navegador se cierra a mitad del proceso lo abre de nuevo desde
    cero (portal -> credenciales -> SIAC) y el lote retoma desde el último RUC del CSV.
    Se rinde tras MAX_REINICIOS_SIN_AVANCE reinicios seguidos sin guardar ningún RUC nuevo."""
    sin_avance = 0
    while True:
        avance_antes = ultimo_ruc_csv()
        try:
            return flujo_portal(url, opcion)
        except Exception as e:
            if not isinstance(e, NavegadorCerrado) and _driver_vivo(_driver_activo):
                raise                         # error con el navegador abierto: no lo tapo
            sin_avance = 0 if ultimo_ruc_csv() != avance_antes else sin_avance + 1
            if sin_avance >= MAX_REINICIOS_SIN_AVANCE:
                log.error("[reinicio] %d reinicios seguidos sin avanzar; me rindo.", sin_avance)
                raise
            log.warning("[reinicio] %s", e if isinstance(e, NavegadorCerrado)
                        else f"El navegador se cerró ({type(e).__name__}: {e})"[:300])
            log.warning("[reinicio] Abro el navegador de nuevo y retomo desde el CSV "
                        "(reinicios seguidos sin avance: %d/%d)...", sin_avance, MAX_REINICIOS_SIN_AVANCE)
            _descartar_navegador()


# --------------------------- cierre ---------------------------
def cerrar_todo(driver):
    """Cierra el navegador. La VPN no se toca: la maneja una persona a mano."""
    if driver is not None and not MANTENER_ABIERTO:
        try:
            driver.quit()
            log.info("[fin] Navegador cerrado.")
        except Exception as e:
            log.warning("[fin] No pude cerrar el navegador: %s", e)


# ---------------------------- main ----------------------------
if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--solo-selenium", action="store_true", help=argparse.SUPPRESS)  # ya es el comportamiento por defecto
    ap.add_argument("--guardar-password", action="store_true",
                    help="Guardar la contraseña en el Administrador de credenciales de Windows")
    ap.add_argument("--guardar-password-portal", action="store_true",
                    help="Guardar la contraseña del portal de distribuidores")
    ap.add_argument("--verificar-password", action="store_true",
                    help="Ver qué contraseña hay guardada y probarla contra el portal")
    ap.add_argument("--metricas", action="store_true",
                    help="Solo mostrar el resumen del CSV (válidos / vacíos / bloqueados)")
    args = ap.parse_args()

    if args.metricas:
        metricas_csv(max_listado=10**6)       # aquí sí lista todos los RUCs
        sys.exit(0)

    codigo_salida = 0
    try:
        if args.verificar_password:
            verificar_password()
            sys.exit(0)

        if args.guardar_password:
            guardar_password()
            sys.exit(0)

        if args.guardar_password_portal:
            guardar_password_portal()
            sys.exit(0)

        # la VPN la conecta una persona a mano antes de ejecutar el script
        flujo_con_reinicios()
        log.info("[fin] Proceso terminado; cerrando todo...")

    except KeyboardInterrupt:
        log.warning("[fin] Interrumpido por el usuario; cerrando todo...")
        codigo_salida = 1
    except Exception as e:
        log.exception("Error: %s", e)
        codigo_salida = 1
    finally:
        cerrar_todo(_driver_activo)
    sys.exit(codigo_salida)


